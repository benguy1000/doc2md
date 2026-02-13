"""DOCX to Markdown converter using python-docx."""

from __future__ import annotations

import base64
import logging
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from doc2md.models import ConversionResult
from doc2md.utils.file_handler import (
    cleanup_temp_dir,
    resolve_output_path,
    resolve_source_file,
    write_markdown,
)
from doc2md.utils.markdown import (
    clean_text,
    count_words,
    format_table,
    generate_frontmatter,
)
from doc2md.utils.metadata import build_metadata

logger = logging.getLogger(__name__)


def _run_to_markdown(run) -> str:
    """Convert a single run to markdown text with formatting."""
    text = run.text
    if not text:
        return ""

    if run.bold:
        text = f"**{text}**"
    if run.italic:
        text = f"*{text}*"
    if run.font.strike:
        text = f"~~{text}~~"
    if run.font.name and "courier" in run.font.name.lower():
        text = f"`{text}`"

    return text


def _paragraph_to_markdown(para: Paragraph, list_counters: dict) -> str:
    """Convert a paragraph to markdown."""
    style_name = para.style.name if para.style else ""

    # Assemble inline text
    text_parts = []
    for run in para.runs:
        text_parts.append(_run_to_markdown(run))
    text = "".join(text_parts).strip()

    if not text:
        return ""

    # Extract hyperlinks from XML
    hyperlinks = para._element.findall(qn("w:hyperlink"))
    for hl in hyperlinks:
        hl_runs = hl.findall(qn("w:r"))
        hl_text = "".join(
            r.find(qn("w:t")).text for r in hl_runs if r.find(qn("w:t")) is not None
        )
        rid = hl.get(qn("r:id"))
        if rid and rid in para.part.rels:
            url = para.part.rels[rid].target_ref
            text = text.replace(hl_text, f"[{hl_text}]({url})", 1)

    # Headings
    if style_name.startswith("Heading"):
        match = re.search(r"(\d+)", style_name)
        level = int(match.group(1)) if match else 1
        level = min(level, 6)
        return f"{'#' * level} {text}"

    # Lists
    numPr = para._element.find(qn("w:pPr"))
    if numPr is not None:
        numPr_el = numPr.find(qn("w:numPr"))
        if numPr_el is not None:
            ilvl_el = numPr_el.find(qn("w:ilvl"))
            numId_el = numPr_el.find(qn("w:numId"))
            indent_level = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
            num_id = numId_el.get(qn("w:val")) if numId_el is not None else "0"
            indent = "  " * indent_level

            # Determine if ordered or unordered
            if "List Number" in style_name or "List Bullet" not in style_name:
                key = f"{num_id}_{indent_level}"
                list_counters[key] = list_counters.get(key, 0) + 1
                return f"{indent}{list_counters[key]}. {text}"
            else:
                return f"{indent}- {text}"

    if "List Bullet" in style_name:
        return f"- {text}"
    if "List Number" in style_name:
        key = f"style_{style_name}"
        list_counters[key] = list_counters.get(key, 0) + 1
        return f"{list_counters[key]}. {text}"

    return text


def _table_to_markdown(table: Table) -> str:
    """Convert a DOCX table to markdown."""
    if not table.rows:
        return ""

    headers = [clean_text(cell.text.strip()) for cell in table.rows[0].cells]
    rows = []
    for row in table.rows[1:]:
        rows.append([clean_text(cell.text.strip()) for cell in row.cells])

    return format_table(headers, rows)


def _extract_images(doc: Document) -> tuple[list[str], int]:
    """Extract embedded images and return (placeholders, count)."""
    placeholders = []
    count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            count += 1
            try:
                image_data = rel.target_part.blob
                ext = Path(rel.target_ref).suffix.lstrip(".")
                if ext in ("png", "jpg", "jpeg", "gif", "bmp", "svg"):
                    b64 = base64.b64encode(image_data).decode("ascii")
                    placeholders.append(
                        f"![embedded image {count}](data:image/{ext};base64,{b64})"
                    )
                else:
                    placeholders.append(f"![embedded image {count}](image_{count}.{ext})")
            except Exception:
                placeholders.append(f"![embedded image {count}](image_{count}.png)")

    return placeholders, count


def _extract_comments(doc: Document) -> list[str]:
    """Extract comments from the document."""
    comments = []
    try:
        comments_part = doc.part.package.part_related_by("/word/comments.xml")
        if comments_part:
            from lxml import etree

            tree = etree.fromstring(comments_part.blob)
            for comment in tree.findall(qn("w:comment")):
                author = comment.get(qn("w:author"), "Unknown")
                texts = comment.findall(".//" + qn("w:t"))
                text = "".join(t.text for t in texts if t.text)
                if text:
                    comments.append(f"<!-- Comment ({author}): {text} -->")
    except Exception:
        pass
    return comments


def _extract_footnotes(doc: Document) -> dict[str, str]:
    """Extract footnotes from the document."""
    footnotes = {}
    try:
        fn_part = doc.part.package.part_related_by("/word/footnotes.xml")
        if fn_part:
            from lxml import etree

            tree = etree.fromstring(fn_part.blob)
            for fn in tree.findall(qn("w:footnote")):
                fn_id = fn.get(qn("w:id"))
                if fn_id and int(fn_id) > 0:  # Skip separator footnotes
                    texts = fn.findall(".//" + qn("w:t"))
                    text = "".join(t.text for t in texts if t.text)
                    if text:
                        footnotes[fn_id] = text
    except Exception:
        pass
    return footnotes


def convert_docx_to_markdown(
    file_path: str | None = None,
    base64_content: str | None = None,
    file_name: str | None = None,
    output_dir: str | None = None,
    output_file_name: str | None = None,
    overwrite: bool = False,
) -> ConversionResult:
    """Convert a DOCX file to Markdown and write to disk."""
    temp_dir = None
    try:
        source_path, source_name, temp_dir = resolve_source_file(
            file_path, base64_content, file_name
        )
        is_base64 = base64_content is not None

        doc = Document(str(source_path))
        warnings: list[str] = []
        sections: list[str] = []
        list_counters: dict = {}

        # Process document body - iterate over block-level elements in order
        body = doc.element.body
        for element in body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag
            if tag == "p":
                para = Paragraph(element, doc)
                md = _paragraph_to_markdown(para, list_counters)
                if md:
                    sections.append(md)
                else:
                    # Preserve empty paragraphs as blank lines (paragraph breaks)
                    if sections and sections[-1] != "":
                        sections.append("")
            elif tag == "tbl":
                table = Table(element, doc)
                md = _table_to_markdown(table)
                if md:
                    sections.append(md)

        # Extract images
        img_placeholders, img_count = _extract_images(doc)
        if img_placeholders:
            sections.append("\n## Embedded Images\n")
            sections.extend(img_placeholders)

        # Extract comments
        comments = _extract_comments(doc)
        if comments:
            sections.append("")
            sections.extend(comments)

        # Extract footnotes
        footnotes = _extract_footnotes(doc)
        if footnotes:
            sections.append("\n## Footnotes\n")
            for fn_id, fn_text in footnotes.items():
                sections.append(f"[^{fn_id}]: {fn_text}")

        body_text = "\n\n".join(s for s in sections if s is not None)
        # Collapse multiple blank lines
        while "\n\n\n\n" in body_text:
            body_text = body_text.replace("\n\n\n\n", "\n\n\n")

        word_count = count_words(body_text)
        page_count = len(doc.sections)

        frontmatter = generate_frontmatter(
            source=source_name,
            fmt="docx",
            pages=page_count,
            word_count=word_count,
            warnings=warnings if warnings else None,
        )

        md_content = frontmatter + "\n" + body_text + "\n"

        output_path = resolve_output_path(
            source_name, source_path, output_dir, output_file_name, overwrite, is_base64
        )
        write_markdown(output_path, md_content)

        metadata = build_metadata(
            source_format="docx",
            page_count=page_count,
            word_count=word_count,
            has_images=img_count > 0,
            image_count=img_count,
            warnings=warnings,
        )

        return ConversionResult(
            success=True,
            output_path=str(output_path),
            file_name=output_path.name,
            source_file=source_name,
            metadata=metadata,
        )

    except Exception as e:
        logger.error("DOCX conversion failed: %s", e)
        return ConversionResult(
            success=False,
            source_file=file_name or (file_path or "unknown"),
            error=str(e),
        )
    finally:
        cleanup_temp_dir(temp_dir)
